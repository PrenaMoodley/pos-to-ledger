import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from pos_ledger.business_central import BCConfig, BusinessCentral, BusinessCentralError
from pos_ledger.journal import AccountMap, JournalError, build_journal
from pos_ledger.pos_export import SalesLine, TenderLine, read_sales, read_tenders
from pos_ledger.reports import load_chart, trial_balance, vat_summary, write_workbook

ROOT = Path(__file__).resolve().parents[1]
DAY1, DAY2 = date(2026, 10, 11), date(2026, 10, 12)


@pytest.fixture
def accounts():
    return AccountMap.load(ROOT / "config" / "accounts.json")


@pytest.fixture
def sample():
    return read_sales(ROOT / "sample_data" / "sales.csv"), read_tenders(ROOT / "sample_data" / "tenders.csv")


def totals(lines):
    return sum(l.debit for l in lines), sum(l.credit for l in lines)


def by_account(lines):
    return {l.account: (l.debit, l.credit) for l in lines}


def test_day_balances_and_splits_vat(sample, accounts):
    lines = build_journal(DAY1, *sample, accounts)
    debit, credit = totals(lines)
    assert debit == credit == Decimal("40920.50")
    acc = by_account(lines)
    # 40 920.50 incl. VAT -> VAT is 15/115 of it.
    assert acc["2200"] == (Decimal("0"), Decimal("5337.46"))
    # (18 450 - 450) / 1.15 = 15 652.17, plus the 1c rounding remainder, which goes
    # to the largest sales line so that sales + VAT equals the tenders exactly.
    assert acc["4100"][1] == Decimal("15652.18")
    assert sum(acc[a][1] for a in ("4100", "4110", "4200", "4210", "4300")) == Decimal("35583.04")
    assert "6900" not in acc  # tenders matched sales exactly


def test_small_difference_goes_to_over_short(sample, accounts):
    lines = build_journal(DAY2, *sample, accounts)
    assert by_account(lines)["6900"] == (Decimal("20.00"), Decimal("0"))
    debit, credit = totals(lines)
    assert debit == credit


def test_large_difference_stops_the_journal(accounts):
    sales = [SalesLine(DAY1, "Restaurant", "Food", Decimal("1000"), Decimal("0"))]
    tenders = [TenderLine(DAY1, "Card", Decimal("800"))]
    with pytest.raises(JournalError, match="over/short limit"):
        build_journal(DAY1, sales, tenders, accounts)


def test_unmapped_revenue_centre_is_reported(accounts):
    sales = [SalesLine(DAY1, "Spa", "Treatments", Decimal("500"), Decimal("0"))]
    with pytest.raises(JournalError, match="Spa\\|Treatments"):
        build_journal(DAY1, sales, [TenderLine(DAY1, "Card", Decimal("500"))], accounts)


def test_vat_rounding_is_exact_across_many_lines(accounts):
    # Five lines of R33.33 each: rounding each line separately would drift.
    sales = [SalesLine(DAY1, "Restaurant", "Food", Decimal("33.33"), Decimal("0"))] * 5
    lines = build_journal(DAY1, sales, [TenderLine(DAY1, "Cash", Decimal("166.65"))], accounts)
    acc = by_account(lines)
    assert acc["2200"][1] + acc["4100"][1] == Decimal("166.65")
    assert acc["2200"][1] == Decimal("21.74")


def test_trial_balance_and_workbook(sample, accounts, tmp_path):
    chart = load_chart(ROOT / "config" / "chart_of_accounts.csv")
    lines = build_journal(DAY1, *sample, accounts) + build_journal(DAY2, *sample, accounts)
    tb = trial_balance(lines, chart, opening={"1010": Decimal("2000.00"), "2200": Decimal("-2000.00")})
    assert sum(r["debit"] for r in tb) == sum(r["credit"] for r in tb)

    vat = vat_summary(lines, accounts.vat_output, set(accounts.sales.values()))
    assert vat["standard_rated_sales_incl_vat"] == Decimal("66440.50")  # 40 920.50 + 25 520.00

    wb = load_workbook(write_workbook(lines, tb, vat, tmp_path / "day.xlsx"))
    assert wb["Journal"].cell(wb["Journal"].max_row, 5).value.startswith("=IF(")
    assert wb["VAT"]["B5"].value == "=ROUND(B2*15/115,2)"


# ---------------------------------------------------------------- Business Central


def bc_with(handler):
    config = BCConfig("tenant", "client", "secret", "Sandbox", "company-1", "POS")
    return BusinessCentral(config, http=httpx.Client(transport=httpx.MockTransport(handler)))


def test_business_central_adds_lines_once(sample, accounts):
    posted = []

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "oauth2/v2.0/token" in url:
            assert b"client_credentials" in request.content
            return httpx.Response(200, json={"access_token": "t", "expires_in": 3600})
        assert request.headers["Authorization"] == "Bearer t"
        if request.method == "GET" and url.split("?")[0].endswith("/journals"):
            return httpx.Response(200, json={"value": [{"id": "j1", "code": "POS"}]})
        if request.method == "GET" and "journalLines" in url:
            already = [p for p in posted if p["documentNumber"] in url]
            return httpx.Response(200, json={"value": already})
        if request.method == "POST" and url.endswith("/journalLines"):
            posted.append(json.loads(request.content))
            return httpx.Response(201, json={})
        return httpx.Response(404)

    bc = bc_with(handler)
    lines = build_journal(DAY1, *sample, accounts)
    assert bc.add_lines(lines) == len(lines)
    assert sum(Decimal(str(p["amount"])) for p in posted) == 0  # debits positive, credits negative
    assert posted[0]["documentNumber"] == "POS-20261011"

    with pytest.raises(BusinessCentralError, match="already in the journal"):
        bc.add_lines(lines)


def test_business_central_error_is_raised():
    def handler(request):
        if "token" in str(request.url):
            return httpx.Response(401, text="invalid_client")
        return httpx.Response(200)

    with pytest.raises(BusinessCentralError, match="Token request failed"):
        bc_with(handler).journal_id()
