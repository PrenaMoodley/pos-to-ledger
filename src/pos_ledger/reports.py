"""Trial balance and VAT summary from journal lines, plus the Excel workbook."""

from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .journal import JournalLine

ZAR = '"R" #,##0.00'


@dataclass(frozen=True)
class Account:
    number: str
    name: str
    category: str  # Asset, Liability, Equity, Income, Expense


def load_chart(path: Path) -> dict[str, Account]:
    with open(path, newline="", encoding="utf-8") as f:
        return {r["number"]: Account(r["number"], r["name"], r["category"]) for r in csv.DictReader(f)}


def trial_balance(
    lines: list[JournalLine], chart: dict[str, Account], opening: dict[str, Decimal] | None = None
) -> list[dict]:
    """One row per account: opening, movement, closing, shown as debit or credit."""
    opening = opening or {}
    movement: dict[str, Decimal] = defaultdict(Decimal)
    for line in lines:
        if line.account not in chart:
            raise KeyError(f"Account {line.account} is not in the chart of accounts")
        movement[line.account] += line.amount
    rows = []
    for number in sorted(set(movement) | set(opening)):
        closing = opening.get(number, Decimal("0")) + movement[number]
        rows.append(
            {
                "account": number,
                "name": chart[number].name,
                "category": chart[number].category,
                "opening": opening.get(number, Decimal("0")),
                "movement": movement[number],
                "debit": closing if closing > 0 else Decimal("0"),
                "credit": -closing if closing < 0 else Decimal("0"),
            }
        )
    return rows


def vat_summary(lines: list[JournalLine], vat_account: str, sales_accounts: set[str]) -> dict:
    """Output VAT for the period and the sales it relates to (supports the VAT201 return)."""
    sales_excl = sum((l.credit - l.debit for l in lines if l.account in sales_accounts), Decimal("0"))
    output_vat = sum((l.credit - l.debit for l in lines if l.account == vat_account), Decimal("0"))
    return {
        "standard_rated_sales_excl_vat": sales_excl,
        "standard_rated_sales_incl_vat": sales_excl + output_vat,
        "output_vat": output_vat,
    }


def write_workbook(lines: list[JournalLine], tb: list[dict], vat: dict, path: Path) -> Path:
    wb = Workbook()
    js = wb.active
    js.title = "Journal"
    js.append(["Date", "Document", "Account", "Description", "Debit", "Credit"])
    for l in lines:
        js.append([l.posting_date, l.document_no, l.account, l.description, l.debit, l.credit])
    last = js.max_row
    js.append(["", "", "", "Totals", f"=SUM(E2:E{last})", f"=SUM(F2:F{last})"])
    js.append(["", "", "", "Balanced?", f'=IF(E{last + 1}=F{last + 1},"Yes","NO")'])

    ts = wb.create_sheet("Trial balance")
    ts.append(["Account", "Name", "Category", "Debit", "Credit"])
    for r in tb:
        ts.append([r["account"], r["name"], r["category"], r["debit"], r["credit"]])
    last_tb = ts.max_row
    ts.append(["", "", "Totals", f"=SUM(D2:D{last_tb})", f"=SUM(E2:E{last_tb})"])
    ts.append(["", "", "Balanced?", f'=IF(ROUND(D{last_tb + 1}-E{last_tb + 1},2)=0,"Yes","NO")'])

    vs = wb.create_sheet("VAT")
    vs.append(["Measure", "Amount"])
    vs.append(["Standard-rated sales incl. VAT", vat["standard_rated_sales_incl_vat"]])
    vs.append(["Standard-rated sales excl. VAT", vat["standard_rated_sales_excl_vat"]])
    vs.append(["Output VAT (journal)", vat["output_vat"]])
    vs.append(["Output VAT (15/115 of incl.)", "=ROUND(B2*15/115,2)"])
    vs.append(["Check", '=IF(ABS(B4-B5)<=0.01,"OK","INVESTIGATE")'])

    for ws in (js, ts, vs):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, Decimal) or (
                    isinstance(cell.value, str) and cell.value.startswith("=SUM")
                ):
                    cell.number_format = ZAR
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
